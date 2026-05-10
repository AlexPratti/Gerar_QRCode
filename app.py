import streamlit as st
import fitz  # PyMuPDF
import re
from supabase import create_client
import pandas as pd
import io
from docx import Document # Adicione no requirements.txt

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(page_title="Hinário Litúrgico", layout="wide")

# Conexão Supabase (Mantendo o que já está certo)
try:
    supabase = create_client(st.secrets["URL_SUPABASE"], st.secrets["KEY_SUPABASE"])
except Exception as e:
    st.error(f"Erro de conexão: {e}")
    st.stop()

CATEGORIAS_ALVO = ["ORANTES", "INICIAIS E FINAIS", "PERDÃO", "GLÓRIA", "DEUS NOS FALA", "SALMO", "ACLAMAÇÃO", "OFERTÓRIO", "LOUVOR", "SANTO", "CORDEIRO", "PAZ", "COMUNHÃO", "BÍBLIA", "CRUZ", "LADAINHAS – SEQUÊNCIAS - PROCLAMAÇÕES", "MARIA", "HINOS DIVERSOS", "PRECES"]

# --- FUNÇÕES DE PROCESSAMENTO PDF ---
def process_pdf_fitz(file_bytes):
    data = []
    current_cat = "Sem Categoria"
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for i in range(len(doc)):
            page = doc[i]
            text = page.get_text()
            if not text: continue
            for line in text.split('\n'):
                t_limpo = line.strip()
                if t_limpo.upper() in CATEGORIAS_ALVO:
                    current_cat = t_limpo.upper()
                elif re.match(r'^\d+\.', t_limpo) and t_limpo == t_limpo.upper():
                    data.append({"n1": current_cat, "n2": t_limpo, "pag": i + 1})
        doc.close()
        return data
    except Exception as e:
        st.error(f"Erro no processamento do PDF: {e}")
        return []

def save_to_db(data, table_cat, table_cont):
    supabase.table(table_cont).delete().neq("id", 0).execute()
    supabase.table(table_cat).delete().neq("id", 0).execute()
    for cat in CATEGORIAS_ALVO:
        res = supabase.table(table_cat).insert({"nome_nivel1": cat}).execute()
        if res.data:
            res_data = res.data if isinstance(res.data, list) else res.data
            cat_id = res_data['id']
            itens = [{"categoria_id": cat_id, "nome_nivel2": item['n2'], "texto_completo": str(item['pag'])} for item in data if item['n1'] == cat]
            if itens: supabase.table(table_cont).insert(itens).execute()

# --- FUNÇÃO PARA LIMPAR CIFRAS (NOVO) ---
def limpar_cifras_docx(file):
    doc = Document(file)
    novo_doc = Document()
    padrao_cifras = r'\b([A-G][b#]?(m|maj|min|7|sus|4|6|9|dim|aug)*)\b'
    
    for p in doc.paragraphs:
        texto = p.text
        limpo = texto.strip()
        if not limpo:
            novo_doc.add_paragraph("")
            continue
        
        acordes = re.findall(padrao_cifras, limpo)
        letras_min = len(re.findall(r'[a-z]', limpo))
        
        # Se tem acordes mas quase não tem letras minúsculas, é linha de cifra
        if len(acordes) > 0 and letras_min < 2:
            continue
        
        novo_doc.add_paragraph(texto)
    
    target = io.BytesIO()
    novo_doc.save(target)
    return target.getvalue()

# --- INTERFACE ---
tab_cifras, tab_letras, tab_up_cifras, tab_up_letras, tab_util = st.tabs([
    "🎸 Hinos com Cifras", "📖 Hinos (Letras)", "⚙️ Upload Cifras", "⚙️ Upload Letras", "🛠️ Limpar Cifras"
])

def render_hino_interface(bucket, file_path, table_cat, table_cont, key_suffix):
    try:
        res_cat = supabase.table(table_cat).select("*").order("nome_nivel1").execute()
        if res_cat.data:
            df_cat = pd.DataFrame(res_cat.data)
            c1, c2 = st.columns(2)
            with c1:
                sel_cat = st.selectbox("Categoria", df_cat['nome_nivel1'], key=f"cat_{key_suffix}")
                cat_id = df_cat[df_cat['nome_nivel1'] == sel_cat]['id'].iloc
            
            hinos_res = supabase.table(table_cont).select("*").eq("categoria_id", int(cat_id)).execute().data
            if hinos_res:
                hinos_ord = sorted(hinos_res, key=lambda x: int(re.search(r'\d+', x['nome_nivel2']).group()) if re.search(r'\d+', x['nome_nivel2']) else 0)
                with c2:
                    sel_hino = st.selectbox("Hino", [h['nome_nivel2'] for h in hinos_ord], key=f"hino_{key_suffix}")
                
                pdf_res = supabase.storage.from_(bucket).download(file_path)
                if pdf_res:
                    hino_obj = next(h for h in hinos_res if h['nome_nivel2'] == sel_hino)
                    p_num = int(hino_obj['texto_completo']) - 1
                    doc = fitz.open(stream=pdf_res, filetype="pdf")
                    page = doc[p_num]
                    text_instances = page.search_for(sel_hino)
                    y_ini = text_instances.y0 if text_instances else 0
                    y_fim = page.rect.height
                    blocks = page.get_text("blocks")
                    for b in blocks:
                        if b[1] > y_ini + 10:
                            txt_block = b[4].strip()
                            if re.match(r'^\d+\.', txt_block) or txt_block.upper() in CATEGORIAS_ALVO:
                                y_fim = b[1]; break
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=fitz.Rect(0, max(0, y_ini-15), page.rect.width, y_fim))
                    st.divider()
                    st.image(pix.tobytes("png"), use_container_width=True)
                    doc.close()
    except Exception as e: st.error(f"Erro: {e}")

with tab_cifras: render_hino_interface("hinarios", "hinario_atual.pdf", "hinos_categorias", "hinos_conteudos", "cifras")
with tab_letras: render_hino_interface("letras", "hinario_letras.pdf", "hinos_categorias_letras", "hinos_conteudos_letras", "letras")

# --- ABAS DE UPLOAD (MANTIDAS) ---
with tab_up_cifras:
    st.subheader("Configuração Cifras")
    n_c = st.file_uploader("PDF Cifras", type="pdf", key="f_c")
    if st.button("🚀 Atualizar Cifras") and n_c:
        b = n_c.read()
        supabase.storage.from_("hinarios").upload(path="hinario_atual.pdf", file=b, file_options={"x-upsert": "true", "content-type": "application/pdf"})
        d = process_pdf_fitz(b)
        if d: save_to_db(d, "hinos_categorias", "hinos_conteudos"); st.success("OK!"); st.rerun()

with tab_up_letras:
    st.subheader("Configuração Letras")
    n_l = st.file_uploader("PDF Letras", type="pdf", key="f_l")
    if st.button("🚀 Atualizar Letras") and n_l:
        b = n_l.read()
        supabase.storage.from_("letras").upload(path="hinario_letras.pdf", file=b, file_options={"x-upsert": "true", "content-type": "application/pdf"})
        d = process_pdf_fitz(b)
        if d: save_to_db(d, "hinos_categorias_letras", "hinos_conteudos_letras"); st.success("OK!"); st.rerun()

# --- ABA 5: UTILITÁRIO DE LIMPEZA ---
with tab_util:
    st.subheader("Remover Cifras de DOCX")
    st.info("Suba o arquivo .docx com cifras e baixe a versão limpa para depois salvar como PDF.")
    arquivo_docx = st.file_uploader("Selecione o arquivo DOCX", type="docx")
    if arquivo_docx:
        if st.button("✨ Limpar Documento"):
            with st.spinner("Removendo cifras..."):
                resultado_bytes = limpar_cifras_docx(arquivo_docx)
                st.success("Documento processado!")
                st.download_button(
                    label="📥 Baixar DOCX Sem Cifras",
                    data=resultado_bytes,
                    file_name="LITURGICOS_SEM_CIFRAS.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
